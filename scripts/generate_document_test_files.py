"""Generate document test files under test_files/documents/ for DataWise feature testing."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "test_files" / "documents"
OUT.mkdir(parents=True, exist_ok=True)


def write_txt() -> None:
    text = """DataWise Sample Report

Invoice No: INV-2024-778
Customer: Alice Smith
Email: alice.smith@example.com
Phone: +1-555-010-2299
Amount: $1,250.00
Date: 2024-03-15
Website: https://example.com/orders/778

This report contains sample customer notes for testing document extraction.
Revenue increased in March. Please find all mentions of Revenue carefully.
Contact support at help@datawise.local if needed.
"""
    (OUT / "sample_report.txt").write_text(text, encoding="utf-8")


def write_md() -> None:
    md = """# Quarterly Overview

## Revenue
Revenue grew by 12% this quarter.

## Customers
- Alice Smith (alice.smith@example.com)
- Bob Jones (bob.jones@example.com)

## Links
See https://example.com/docs for details.

### Notes
Phone: 555-222-8899
"""
    (OUT / "sample_overview.md").write_text(md, encoding="utf-8")


def write_docx() -> None:
    doc = Document()
    doc.add_heading("Customer Invoice Pack", level=1)
    doc.add_paragraph("Invoice No: INV-1001")
    doc.add_paragraph("Customer Name: Diana Prince")
    doc.add_paragraph("Email: diana.prince@example.com")
    doc.add_paragraph("Phone: +44 20 7946 0958")
    doc.add_paragraph("Total Amount: USD 980.50")
    doc.add_paragraph("Due Date: 15/04/2024")
    doc.add_heading("Line Items", level=2)
    table = doc.add_table(rows=4, cols=3)
    headers = ["Item", "Qty", "Price"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows = [
        ("Widget A", "2", "100.00"),
        ("Widget B", "1", "480.50"),
        ("Service Fee", "1", "300.00"),
    ]
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            table.rows[r_i].cells[c_i].text = val
    doc.add_paragraph("Visit https://billing.example.com for payment.")
    doc.save(str(OUT / "sample_invoice.docx"))


def write_pdf() -> None:
    path = OUT / "sample_report.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 720, "DataWise PDF Report")
    c.setFont("Helvetica", 11)
    lines = [
        "Invoice No: INV-2024-778",
        "Customer: Alice Smith",
        "Email: alice.smith@example.com",
        "Phone: +1-555-010-2299",
        "Amount: $1,250.00",
        "Date: March 15, 2024",
        "Website: https://example.com/orders/778",
        "",
        "Revenue increased in March across all regions.",
        "Please review mentions of Revenue in this document.",
        "Contact help@datawise.local for support.",
    ]
    y = 690
    for line in lines:
        c.drawString(72, y, line)
        y -= 18
    c.showPage()
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "Page 2 — Details")
    c.setFont("Helvetica", 11)
    c.drawString(72, 690, "Second page content for split_pages testing.")
    c.drawString(72, 670, "Region: North | Score: 88")
    # Simple table-like text block (PDF table extraction is best-effort)
    c.drawString(72, 640, "Item | Qty | Price")
    c.drawString(72, 622, "Alpha | 3 | 40.00")
    c.drawString(72, 604, "Beta | 1 | 75.00")
    c.save()


def write_readme() -> None:
    (OUT / "README.md").write_text(
        """# Document test files

Use these files in the DataWise UI or API to exercise document features.

| File | Format | Good for |
| :--- | :--- | :--- |
| `sample_report.pdf` | PDF | text, entities, key-values, mentions, links, pages, metadata |
| `sample_invoice.docx` | DOCX | text, tables→CSV, headings, entities, key-values |
| `sample_report.txt` | TXT | text, entities, redact, summary, language |
| `sample_overview.md` | MD | headings, links, entities, summary |

Example prompts are listed in `USAGE_GUIDE.md` (repo root).
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    write_txt()
    write_md()
    write_docx()
    write_pdf()
    write_readme()
    print(f"Wrote document test files to {OUT}")
